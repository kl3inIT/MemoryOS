from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor


SOURCE = Path(
    r"D:\FPT_Curriculum\Semester_9\MemoryOS\tmp\pdfs\agentic-harness-casan\source-copy.docx"
)
OUTPUT = Path(
    r"D:\FPT_Curriculum\Semester_9\MemoryOS\tmp\pdfs\agentic-harness-casan\styled-for-pdf.docx"
)

FONT = "Segoe UI"
NAVY = "1F4E79"
BLUE = "2F75B5"
BODY = "263238"
MUTED = "64748B"
LIGHT_BLUE = "D9E5F2"
TABLE_BORDER = "B7C9DC"
TABLE_ALT = "F4F8FC"


def set_run_font(run, size, color=BODY, bold=False, italic=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        r_fonts.set(qn(f"w:{attr}"), FONT)
    lang = r_pr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        r_pr.append(lang)
    lang.set(qn("w:val"), "vi-VN")


def set_paragraph_border(paragraph, edge, color, size=6, space=4):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    border = p_bdr.find(qn(f"w:{edge}"))
    if border is None:
        border = OxmlElement(f"w:{edge}")
        p_bdr.append(border)
    border.set(qn("w:val"), "single")
    border.set(qn("w:sz"), str(size))
    border.set(qn("w:space"), str(space))
    border.set(qn("w:color"), color)


def set_bottom_border(paragraph, color, size=6, space=4):
    set_paragraph_border(paragraph, "bottom", color, size, space)


def clear_paragraph_borders(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is not None:
        p_pr.remove(p_bdr)


def remove_empty_paragraphs(document):
    for paragraph in list(document.paragraphs):
        if not paragraph.text.strip():
            parent = paragraph._element.getparent()
            parent.remove(paragraph._element)


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_cell_fill(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, left=110, bottom=80, right=110):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        element = tc_mar.find(qn(f"w:{side}"))
        if element is None:
            element = OxmlElement(f"w:{side}")
            tc_mar.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def set_table_borders(table, color=TABLE_BORDER, size=5):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), str(size))
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), color)


def set_table_widths(table, widths_cm):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    total_twips = int(sum(widths_cm) / 2.54 * 1440)
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.insert(0, tbl_w)
    tbl_w.set(qn("w:w"), str(total_twips))
    tbl_w.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    grid_cols = list(grid)
    for index, width_cm in enumerate(widths_cm):
        width_twips = int(width_cm / 2.54 * 1440)
        if index < len(grid_cols):
            grid_cols[index].set(qn("w:w"), str(width_twips))
        for row in table.rows:
            cell = row.cells[index]
            cell.width = Cm(width_cm)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width_twips))
            tc_w.set(qn("w:type"), "dxa")


def add_field(paragraph, instruction):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" {instruction} "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    set_run_font(run, 8.5, MUTED)


def format_footer(footer):
    paragraph = footer.paragraphs[0]
    paragraph.clear()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(0)
    set_paragraph_border(paragraph, "top", LIGHT_BLUE, size=4, space=4)
    label = paragraph.add_run("Trang ")
    set_run_font(label, 8.5, MUTED)
    add_field(paragraph, "PAGE")
    middle = paragraph.add_run(" / ")
    set_run_font(middle, 8.5, MUTED)
    add_field(paragraph, "NUMPAGES")


def format_header(header):
    paragraph = header.paragraphs[0]
    paragraph.clear()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run("MEMORY OS  /  AGENTIC AI")
    set_run_font(run, 8.0, MUTED, bold=True)
    set_bottom_border(paragraph, LIGHT_BLUE, size=4, space=3)


def style_document():
    document = Document(SOURCE)
    remove_empty_paragraphs(document)

    section = document.sections[0]
    section.start_type = WD_SECTION_START.NEW_PAGE
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.left_margin = Mm(22)
    section.right_margin = Mm(22)
    section.top_margin = Mm(20)
    section.bottom_margin = Mm(18)
    section.header_distance = Mm(9)
    section.footer_distance = Mm(9)
    section.different_first_page_header_footer = True
    document.settings.odd_and_even_pages_header_footer = False

    format_header(section.header)
    section.first_page_header.paragraphs[0].clear()
    format_footer(section.footer)
    format_footer(section.even_page_footer)
    format_footer(section.first_page_footer)

    for index, paragraph in enumerate(document.paragraphs):
        style_name = paragraph.style.name
        clear_paragraph_borders(paragraph)
        fmt = paragraph.paragraph_format
        fmt.widow_control = True

        if index == 0:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            fmt.left_indent = Pt(0)
            fmt.right_indent = Pt(0)
            fmt.first_line_indent = Pt(0)
            fmt.space_before = Pt(4)
            fmt.space_after = Pt(18)
            fmt.line_spacing = 1.08
            fmt.keep_with_next = True
            for run in paragraph.runs:
                set_run_font(run, 19.0, NAVY, bold=True)
            set_bottom_border(paragraph, BLUE, size=10, space=9)
            continue

        if style_name == "Heading 1":
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            fmt.left_indent = Pt(0)
            fmt.right_indent = Pt(0)
            fmt.first_line_indent = Pt(0)
            fmt.space_before = Pt(12)
            fmt.space_after = Pt(5)
            fmt.line_spacing = 1.05
            fmt.keep_with_next = True
            fmt.keep_together = True
            if paragraph.text.startswith("4. "):
                fmt.page_break_before = True
            for run in paragraph.runs:
                set_run_font(run, 13.5, NAVY, bold=True)
            set_bottom_border(paragraph, LIGHT_BLUE, size=6, space=3)
        elif style_name == "Heading 2":
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            fmt.left_indent = Cm(0.45)
            fmt.right_indent = Pt(0)
            fmt.first_line_indent = Pt(0)
            fmt.space_before = Pt(7)
            fmt.space_after = Pt(3)
            fmt.line_spacing = 1.05
            fmt.keep_with_next = True
            fmt.keep_together = True
            for run in paragraph.runs:
                set_run_font(run, 11.5, BLUE, bold=True)
        elif style_name == "List Bullet":
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            fmt.left_indent = Cm(1.05)
            fmt.first_line_indent = Cm(-0.42)
            fmt.right_indent = Pt(0)
            fmt.space_before = Pt(0)
            fmt.space_after = Pt(1.5)
            fmt.line_spacing = 1.1
            for run in paragraph.runs:
                set_run_font(run, 10.5, BODY)
        else:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            fmt.left_indent = Cm(0.85)
            fmt.right_indent = Pt(0)
            fmt.first_line_indent = Pt(0)
            fmt.space_before = Pt(0)
            fmt.space_after = Pt(4)
            fmt.line_spacing = 1.17
            for run in paragraph.runs:
                set_run_font(run, 10.5, BODY)

    width_sets = [
        [3.6, 13.0],
        [4.5, 12.1],
        [3.2, 4.1, 5.0, 4.3],
        [3.4, 13.2],
    ]
    for table_index, table in enumerate(document.tables):
        set_table_widths(table, width_sets[table_index])
        set_table_borders(table)
        set_repeat_header(table.rows[0])
        for row_index, row in enumerate(table.rows):
            prevent_row_split(row)
            for cell_index, cell in enumerate(row.cells):
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                set_cell_margins(cell)
                if row_index == 0:
                    set_cell_fill(cell, BLUE)
                elif row_index % 2 == 0:
                    set_cell_fill(cell, TABLE_ALT)
                else:
                    set_cell_fill(cell, "FFFFFF")
                for paragraph in cell.paragraphs:
                    paragraph.alignment = (
                        WD_ALIGN_PARAGRAPH.CENTER if row_index == 0 else WD_ALIGN_PARAGRAPH.LEFT
                    )
                    paragraph.paragraph_format.left_indent = Pt(0)
                    paragraph.paragraph_format.right_indent = Pt(0)
                    paragraph.paragraph_format.first_line_indent = Pt(0)
                    paragraph.paragraph_format.space_before = Pt(0)
                    paragraph.paragraph_format.space_after = Pt(1)
                    paragraph.paragraph_format.line_spacing = 1.05
                    paragraph.paragraph_format.keep_together = True
                    for run in paragraph.runs:
                        set_run_font(
                            run,
                            9.4,
                            "FFFFFF" if row_index == 0 else BODY,
                            bold=(row_index == 0 or (cell_index == 0 and row_index > 0)),
                        )

    document.core_properties.title = "Tính Agentic AI của Memory OS qua hai bài toán của Tasco"
    document.core_properties.subject = "Agentic AI, Harness Engineering và CASAN"
    document.core_properties.keywords = "Memory OS, Agentic AI, Harness Engineering, CASAN, Tasco"
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    style_document()
