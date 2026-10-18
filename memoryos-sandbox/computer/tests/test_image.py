"""The computer image's Python stack and office tools, run in the image as a sandbox runs them (MEM-271).

Moved from the interpreter executor's integration tests (MEM-110) when run_python moved to the sandbox: the same
scripts run as the sandbox user from /workspace, without a network. Standard library only.

IMAGE=memoryos-sandbox-computer:<tag> python3 -m unittest discover -s memoryos-sandbox/computer/tests
"""

from __future__ import annotations

import json
import os
import shutil
import subprocess
import tempfile
import unittest
from pathlib import Path

IMAGE = os.environ.get("IMAGE", "")
VIETNAMESE = "Báo cáo doanh thu quý 3 — Hà Nội, Đà Nẵng"
# onnxruntime (via markitdown's magika) logs this at import on Hyper-V hosts such as GitHub runners;
# no setting silences it (microsoft/onnxruntime#27092).
ONNXRUNTIME_DEVICE_DISCOVERY_WARNING = "device_discovery.cc"


@unittest.skipUnless(IMAGE, "set IMAGE to the computer image to test")
class ComputerImageTest(unittest.TestCase):
    def setUp(self) -> None:
        self.workspace = Path(tempfile.mkdtemp(prefix="computer-image-"))
        self.workspace.chmod(0o777)  # the sandbox user (1000) writes into it

    def tearDown(self) -> None:
        # Files the sandbox user wrote are removed from inside, as the runner may not own them.
        subprocess.run(["docker", "run", "--rm", "-v", f"{self.workspace}:/workspace", "--entrypoint", "sh", IMAGE,
                        "-c", "rm -rf /workspace/* /workspace/.[!.]*"], check=False, capture_output=True)
        shutil.rmtree(self.workspace, ignore_errors=True)

    def run_python(self, code: str, *, clean: bool = True) -> subprocess.CompletedProcess[str]:
        run = subprocess.run(
            ["docker", "run", "--rm", "--network", "none", "--user", "1000", "-v", f"{self.workspace}:/workspace",
             "-w", "/workspace", IMAGE, "python3", "-c", code],
            capture_output=True, text=True, encoding="utf-8", timeout=180,
        )
        if clean:
            self.assertEqual(run.returncode, 0, run.stderr)
            stderr = [line for line in run.stderr.splitlines() if ONNXRUNTIME_DEVICE_DISCOVERY_WARNING not in line]
            self.assertEqual(stderr, [], run.stderr)
        return run

    def test_added_packages_and_cli_tools_are_available(self) -> None:
        run = self.run_python("""
import importlib, json, shutil
modules = ['statsmodels', 'pyarrow', 'xlrd', 'xlsxwriter', 'chardet', 'charset_normalizer',
           'tabulate', 'jinja2', 'markdown', 'bs4', 'markitdown', 'pdf2image', 'sympy']
for name in modules:
    importlib.import_module(name)
tools = ['pdftoppm', 'pdftotext', 'qpdf', 'sqlite3', 'unzip', 'zip', 'soffice', 'recalc-xlsx', 'check-docx']
print(json.dumps({'missing_tools': [t for t in tools if shutil.which(t) is None]}))
""")
        self.assertEqual(json.loads(run.stdout), {"missing_tools": []})

    def test_vietnamese_pdf_with_registered_font_round_trips(self) -> None:
        run = self.run_python(f"""
from fpdf import FPDF
from pypdf import PdfReader
pdf = FPDF()
pdf.add_page()
pdf.add_font('DejaVu', fname='/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf')
pdf.set_font('DejaVu', size=14)
pdf.cell(text={VIETNAMESE!r})
pdf.output('report.pdf')
print(PdfReader('report.pdf').pages[0].extract_text().strip())
""")
        self.assertEqual(run.stdout, VIETNAMESE + "\n")
        self.assertTrue((self.workspace / "report.pdf").read_bytes().startswith(b"%PDF-"))

    def test_recalc_xlsx_fills_formula_values_and_reports_errors(self) -> None:
        run = self.run_python("""
import json, subprocess, zipfile, openpyxl, xlsxwriter
from openpyxl.chart import BarChart, Reference
wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Doanh thu'
for row in [['Tháng', 'Doanh thu'], [1, 120], [2, 80], [3, 100]]:
    ws.append(row)
ws['B5'] = '=SUM(B2:B4)'; ws['B6'] = '=B5/0'
chart = BarChart()
data = Reference(ws, min_col=2, min_row=1, max_row=4)
chart.add_data(data, titles_from_data=True)
ws.add_chart(chart, 'D2'); wb.save('báo cáo.xlsx')
x = xlsxwriter.Workbook('cached.xlsx'); s = x.add_worksheet()
s.write_column('A1', [10, 20, 30]); s.write_formula('A4', '=SUM(A1:A3)'); x.close()
before = openpyxl.load_workbook('báo cáo.xlsx', data_only=True)['Doanh thu']['B5'].value
files = ['báo cáo.xlsx', 'cached.xlsx', 'macro.xlsm', 'cached.xlsx']
done = subprocess.run(['recalc-xlsx', *files], capture_output=True, text=True)
reports = [json.loads(line) for line in done.stdout.splitlines()]
book = openpyxl.load_workbook('báo cáo.xlsx', data_only=True)['Doanh thu']
names = zipfile.ZipFile('báo cáo.xlsx').namelist()
print(json.dumps({
    'returncode': done.returncode, 'before': before, 'sum': book['B5'].value,
    'cached': openpyxl.load_workbook('cached.xlsx', data_only=True).active['A4'].value,
    'formula': openpyxl.load_workbook('báo cáo.xlsx')['Doanh thu']['B5'].value,
    'charts': sum(n.startswith('xl/charts/chart') for n in names),
    'reports': reports,
}, ensure_ascii=False))
""", clean=False)
        result = json.loads(run.stdout)
        self.assertIsNone(result["before"])
        self.assertEqual(result["sum"], 300)
        self.assertEqual(result["cached"], 60)
        self.assertEqual(result["formula"], "=SUM(B2:B4)")
        self.assertEqual(result["charts"], 1)
        self.assertEqual(result["returncode"], 1)  # the .xlsm argument is refused
        self.assertEqual(result["reports"][0], {"file": "báo cáo.xlsx", "recalculated": True, "error_count": 1,
                                                "errors": ["Doanh thu!B6: #DIV/0!"]})
        self.assertIs(result["reports"][1]["recalculated"], True)
        self.assertEqual(result["reports"][2], {"file": "macro.xlsm", "recalculated": False,
                                                "error": "only .xlsx files are supported"})
        self.assertEqual(result["reports"][3], {"file": "cached.xlsx", "recalculated": False,
                                                "error": "run separately: same file name as another argument"})

    def test_open_figures_are_captured_as_chart_data_and_png_at_exit(self) -> None:
        self.run_python("""
import matplotlib.pyplot as plt
fig, ax = plt.subplots()
ax.plot(['T1', 'T2', 'T3'], [120, 150, 90], label='Doanh thu')
ax.set_title('Doanh thu quý 3'); ax.set_ylabel('Tỷ đồng (₫)')
plt.figure(); plt.pie([40, 60], labels=['Hà Nội', 'Đà Nẵng'])
saved = plt.figure(); plt.bar(['A'], [1]); saved.savefig('mine.png'); plt.close(saved)
""")
        charts = self.workspace / ".memoryos-charts"
        self.assertTrue((self.workspace / "mine.png").exists())
        self.assertTrue((charts / "chart-1.png").read_bytes().startswith(b"\x89PNG"))
        self.assertTrue((charts / "chart-2.png").exists())
        self.assertFalse((charts / "chart-3.png").exists())  # a closed figure is not captured
        line = json.loads((charts / "chart-1.json").read_text(encoding="utf-8"))
        self.assertEqual(line["type"], "line")
        self.assertEqual(line["title"], "Doanh thu quý 3")
        self.assertEqual(line["y_unit"], "₫")
        self.assertEqual(line["elements"], [{"label": "Doanh thu", "points": [["T1", 120.0], ["T2", 150.0], ["T3", 90.0]]}])
        pie = json.loads((charts / "chart-2.json").read_text(encoding="utf-8"))
        self.assertEqual(pie["type"], "pie")
        self.assertEqual([element["label"] for element in pie["elements"]], ["Hà Nội", "Đà Nẵng"])

    def test_runs_without_pyplot_leave_no_chart_directory(self) -> None:
        self.run_python("print('xin chào')")
        self.assertFalse((self.workspace / ".memoryos-charts").exists())

    def test_check_docx_names_the_formatting_faults_and_passes_a_structured_document(self) -> None:
        run = self.run_python("""
import json, subprocess
from docx import Document
flat = Document()
for line in ['Báo cáo quý 3', 'Doanh thu tăng.', '• Hà Nội', '1. Đà Nẵng', 'Ghi chú\\nhai dòng',
             'Chi phí giảm.', 'Lợi nhuận tăng.', 'Kết luận.']:
    flat.add_paragraph(line)
table = flat.add_table(rows=2, cols=2)
table.rows[1].cells[0].text = '120'
flat.save('phẳng.docx')
good = Document()
good.add_heading('Báo cáo quý 3', level=0)
good.add_heading('Doanh thu', level=1)
good.add_paragraph('Doanh thu tăng 12%.')
good.add_paragraph('Hà Nội', style='List Bullet')
grid = good.add_table(rows=2, cols=2)
grid.style = 'Table Grid'
grid.rows[0].cells[0].text = 'Tháng'; grid.rows[0].cells[1].text = 'Doanh thu'
grid.rows[1].cells[0].text = '9'; grid.rows[1].cells[1].text = '120'
good.save('đẹp.docx')
done = subprocess.run(['check-docx', 'phẳng.docx', 'đẹp.docx', 'thiếu.docx'], capture_output=True, text=True)
print(json.dumps({'returncode': done.returncode,
                  'reports': [json.loads(line) for line in done.stdout.splitlines()]}, ensure_ascii=False))
""", clean=False)
        self.assertEqual(run.returncode, 0, run.stderr)
        result = json.loads(run.stdout)
        flat, good, missing = result["reports"]
        issues = " | ".join(flat["issues"])
        for fault in ("no heading", "typed bullet", "numbers itself", "line break", "no borders", "empty first row"):
            self.assertIn(fault, issues)
        # A document that already uses headings, list styles and a bordered table reports nothing.
        self.assertEqual(good["issue_count"], 0, good)
        self.assertEqual(good["tables"], 1)
        # An unreadable argument is reported, not raised, and only that turns the exit code non-zero.
        self.assertEqual(missing, {"file": "thiếu.docx", "checked": False, "error": "file not found"})
        self.assertEqual(result["returncode"], 1)

    def test_pptx_to_pdf_converts_a_vietnamese_deck_and_refuses_other_files(self) -> None:
        run = self.run_python("""
import json, subprocess
from pptx import Presentation
deck = Presentation()
for number in range(3):
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = f'Doanh thu quý 3 — slide {number + 1}'
    slide.placeholders[1].text = 'Miền Bắc tăng 22,1% • Hà Nội, Đà Nẵng'
deck.save('báo cáo.pptx')
done = subprocess.run(['pptx-to-pdf', 'báo cáo.pptx', 'preview.pdf'], capture_output=True, text=True)
text = subprocess.run(['pdftotext', 'preview.pdf', '-'], capture_output=True, text=True).stdout
refused = subprocess.run(['pptx-to-pdf', 'notes.txt', 'x.pdf'], capture_output=True, text=True)
print(json.dumps({'code': done.returncode, 'report': json.loads(done.stdout), 'text': text,
                  'refused': refused.returncode, 'refusal': json.loads(refused.stdout)}, ensure_ascii=False))
""", clean=False)
        result = json.loads(run.stdout)
        self.assertEqual(result["code"], 0)
        self.assertEqual(result["report"], {"converted": True, "pages": 3})
        self.assertIn("Doanh thu quý 3 — slide 2", result["text"])
        self.assertIn("Miền Bắc tăng 22,1%", result["text"])
        self.assertEqual(result["refused"], 1)
        self.assertEqual(result["refusal"], {"converted": False, "error": "only .pptx files are supported"})

    def test_files_written_to_mnt_data_land_in_the_workspace(self) -> None:
        """ChatGPT-trained models save to /mnt/data; it is the workspace."""
        self.run_python("open('/mnt/data/bao-cao.txt', 'w', encoding='utf-8').write('Q3')")
        self.assertEqual((self.workspace / "bao-cao.txt").read_text(encoding="utf-8"), "Q3")

    def test_the_idioms_the_prompt_teaches_hold_on_the_image(self) -> None:
        # RUN_PYTHON_GUIDANCE in core's ChatPrompts names these versions and idioms. Each check proves one claim, so a
        # relock that changes what the model is told fails here: update the prompt with it.
        run = self.run_python("""
import io, json, sys
import cv2
import numpy as np
import openpyxl
import pandas as pd
checks = {}
checks['python'] = list(sys.version_info[:2])
checks['numpy'] = '.'.join(np.__version__.split('.')[:2])
checks['pandas'] = '.'.join(pd.__version__.split('.')[:2])
checks['opencv'] = int(cv2.__version__.split('.')[0])
text = pd.DataFrame({'a': ['x', 'y']})['a']
checks['str_dtype'] = str(text.dtype) == 'str' and pd.api.types.is_string_dtype(text)
df = pd.DataFrame({'a': [1.0, None]})
df['a'] = df['a'].fillna(0)
checks['assign_back'] = df['a'].tolist() == [1.0, 0.0]
frame = pd.DataFrame({'a': [1, 2]})
try:
    frame['a'][0] = 9
    checks['chained_raises'] = False
except pd.errors.ChainedAssignmentError:
    checks['chained_raises'] = frame['a'].tolist() == [1, 2]
frame.loc[frame['a'] == 1, 'a'] = 9
checks['loc_assign'] = frame['a'].tolist() == [9, 2]
checks['freq_me'] = len(pd.date_range('2024-01-01', periods=3, freq='ME')) == 3
try:
    pd.date_range('2024-01-01', periods=3, freq='M')
    checks['freq_m_rejected'] = False
except ValueError:
    checks['freq_m_rejected'] = True
hourly = pd.Series([1, 2], index=pd.date_range('2024-01-01', periods=2, freq='h'))
checks['resample'] = hourly.resample('D').sum().tolist() == [3]
checks['ffill'] = pd.Series([1.0, None]).ffill().tolist() == [1.0, 1.0]
book = openpyxl.Workbook()
sheet = book.active
sheet.append(['name', 'amount']); sheet.append(['a', 1]); sheet.append(['b', 2])
buffer = io.BytesIO()
book.save(buffer)
read = pd.read_excel(io.BytesIO(buffer.getvalue()))
checks['excel_dtypes'] = [str(read['name'].dtype), str(read['amount'].dtype)] == ['str', 'int64']
pivot = pd.DataFrame({'k': ['a', 'a', 'b'], 'v': [1, 2, 3]}).pivot_table(index='k', values='v', aggfunc='sum')
checks['pivot'] = pivot['v'].tolist() == [3, 3]
try:
    pd.to_datetime(pd.Series(['2024-05-01', '2024-05-02']), dayfirst=True)
    checks['iso_dayfirst_raises'] = False
except ValueError as error:
    checks['iso_dayfirst_raises'] = 'ISO' in str(error)
parsed = pd.to_datetime(pd.Series(['05/03/2024', '13/03/2024']), dayfirst=True)
checks['day_first_text'] = list(parsed.dt.strftime('%Y-%m-%d')) == ['2024-03-05', '2024-03-13']
checks['iso_plain'] = str(pd.to_datetime('2024-05-01').date()) == '2024-05-01'
checks['scalar_repr'] = repr(np.float64(1.5)) == 'np.float64(1.5)'
checks['scalar_item'] = json.dumps([np.float64(1.5).item()]) == '[1.5]'
checks['nan_alias_removed'] = not hasattr(np, 'NaN')
checks['trapezoid'] = float(np.trapezoid([1.0, 1.0])) == 1.0
checks['cv2_encode'] = bool(cv2.imencode('.png', np.zeros((4, 4, 3), np.uint8))[0])
checks['cv2_ml_absent'] = not hasattr(cv2, 'ml')
checks['cv2_haar_absent'] = not hasattr(cv2, 'CascadeClassifier')
print(json.dumps(checks))
""", clean=False)
        self.assertEqual(run.returncode, 0, run.stderr)
        checks = json.loads(run.stdout)
        self.assertEqual(checks.pop("python"), [3, 14])
        self.assertEqual((checks.pop("numpy"), checks.pop("pandas"), checks.pop("opencv")), ("2.5", "3.0", 5))
        self.assertEqual({name: ok for name, ok in checks.items() if not ok}, {})


if __name__ == "__main__":
    unittest.main()
