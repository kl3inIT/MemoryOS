#!/usr/bin/env python3
"""Fill a PE DOCX template from a small JSON manifest."""

from __future__ import annotations

import argparse
import json
from pathlib import Path
from string import Template
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Inches, RGBColor


def all_paragraphs(document: Document):
    yield from document.paragraphs
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs


def substitute_variables(value: Any, variables: dict[str, Any]) -> Any:
    if isinstance(value, str):
        return Template(value).substitute(
            {key: str(item) for key, item in variables.items()}
        )
    if isinstance(value, list):
        return [substitute_variables(item, variables) for item in value]
    if isinstance(value, dict):
        return {
            key: substitute_variables(item, variables)
            for key, item in value.items()
        }
    return value


def replace_text(paragraph, old: str, new: str) -> bool:
    if old not in paragraph.text:
        return False
    combined = paragraph.text.replace(old, new)
    if paragraph.runs:
        paragraph.runs[0].text = combined
        paragraph.runs[0].font.color.rgb = RGBColor(0, 0, 0)
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        run = paragraph.add_run(combined)
        run.font.color.rgb = RGBColor(0, 0, 0)
    return True


def set_text(paragraph, new: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = new
        paragraph.runs[0].font.color.rgb = RGBColor(0, 0, 0)
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        run = paragraph.add_run(new)
        run.font.color.rgb = RGBColor(0, 0, 0)


def replace_image(document: Document, item: dict[str, Any], base: Path) -> None:
    marker = item["marker"]
    matches = [p for p in all_paragraphs(document) if marker in p.text]
    if len(matches) != 1:
        raise ValueError(f"image marker must match exactly once: {marker!r}; got {len(matches)}")
    paragraph = matches[0]
    for run in paragraph.runs:
        run.text = ""
    run = paragraph.runs[0] if paragraph.runs else paragraph.add_run()
    path = (base / item["path"]).resolve()
    width = Inches(float(item.get("width_inches", 6.25)))
    run.add_picture(str(path), width=width)


def replace_table(document: Document, item: dict[str, Any]) -> None:
    index = int(item["index"])
    table = document.tables[index]
    keep_rows = int(item.get("header_rows", 1))
    while len(table.rows) > keep_rows:
        table._tbl.remove(table.rows[-1]._tr)

    for row_values in item["rows"]:
        cells = table.add_row().cells
        if len(row_values) != len(cells):
            raise ValueError(
                f"table {index} expects {len(cells)} cells, got {len(row_values)}"
            )
        for cell, value in zip(cells, row_values):
            cell.text = str(value)

    for row in table.rows:
        row_properties = row._tr.get_or_add_trPr()
        if row_properties.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}cantSplit") is None:
            row_properties.append(OxmlElement("w:cantSplit"))

    if item.get("keep_together"):
        for row in table.rows[:-1]:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.keep_with_next = True


def remove_paragraph(paragraph) -> None:
    element = paragraph._element
    element.getparent().remove(element)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--template", required=True, type=Path)
    parser.add_argument("--manifest", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("--variables", type=Path)
    args = parser.parse_args()

    manifest = json.loads(args.manifest.read_text(encoding="utf-8"))
    if args.variables:
        variables = json.loads(args.variables.read_text(encoding="utf-8"))
        manifest = substitute_variables(manifest, variables)
    document = Document(args.template)

    for item in manifest.get("paragraph_replacements", []):
        if "paragraph_index" in item:
            paragraph = document.paragraphs[int(item["paragraph_index"])]
            expected = item.get("expected_text")
            if expected is not None and paragraph.text != expected:
                raise ValueError(
                    f"paragraph {item['paragraph_index']}: expected {expected!r}, "
                    f"got {paragraph.text!r}"
                )
            set_text(paragraph, item["replace"])
            continue

        matches = 0
        for paragraph in all_paragraphs(document):
            matches += int(replace_text(paragraph, item["find"], item["replace"]))
        expected = int(item.get("expected_matches", 1))
        if matches != expected:
            raise ValueError(
                f"replacement {item['find']!r}: expected {expected}, got {matches}"
            )

    manifest_base = args.manifest.parent
    for item in manifest.get("images", []):
        replace_image(document, item, manifest_base)

    for item in manifest.get("tables", []):
        replace_table(document, item)

    keep_texts = set(manifest.get("keep_with_next_paragraphs", []))
    for paragraph in document.paragraphs:
        if paragraph.text in keep_texts:
            paragraph.paragraph_format.keep_with_next = True

    paragraphs = list(document.paragraphs)
    alignment_map = {
        "LEFT": WD_ALIGN_PARAGRAPH.LEFT,
        "CENTER": WD_ALIGN_PARAGRAPH.CENTER,
        "RIGHT": WD_ALIGN_PARAGRAPH.RIGHT,
        "JUSTIFY": WD_ALIGN_PARAGRAPH.JUSTIFY,
    }
    for item in manifest.get("paragraph_alignments", []):
        paragraphs[int(item["paragraph_index"])].alignment = alignment_map[
            item["alignment"].upper()
        ]

    for index in sorted(manifest.get("remove_paragraph_indices", []), reverse=True):
        remove_paragraph(paragraphs[int(index)])

    args.output.parent.mkdir(parents=True, exist_ok=True)
    document.save(args.output)
    print(args.output)


if __name__ == "__main__":
    main()
