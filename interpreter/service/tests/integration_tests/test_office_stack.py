"""MemoryOS executor additions (MEM-110): packages, CLI tools and Vietnamese PDF output."""

from __future__ import annotations

import json
from typing import Any, cast

from fastapi.testclient import TestClient

from memoryos_interpreter.main import create_app

VIETNAMESE = "Báo cáo doanh thu quý 3 — Hà Nội, Đà Nẵng"

# onnxruntime (via markitdown's magika) logs this at import on Hyper-V hosts such as GitHub runners;
# no setting silences it (microsoft/onnxruntime#27092).
ONNXRUNTIME_DEVICE_DISCOVERY_WARNING = "device_discovery.cc"


def _execute(client: TestClient, code: str) -> dict[str, object]:
    response = client.post("/v1/execute", json={"code": code, "timeout_ms": 30000})
    assert response.status_code == 200
    payload: dict[str, object] = response.json()
    assert payload["exit_code"] == 0, payload
    stderr = [
        line
        for line in str(payload["stderr"]).splitlines()
        if ONNXRUNTIME_DEVICE_DISCOVERY_WARNING not in line
    ]
    assert stderr == [], payload
    return payload


def test_added_packages_and_cli_tools_are_available() -> None:
    client = TestClient(create_app())
    code = """
import importlib, json, shutil
modules = ['statsmodels', 'pyarrow', 'xlrd', 'xlsxwriter', 'chardet', 'charset_normalizer',
           'tabulate', 'jinja2', 'markdown', 'bs4', 'markitdown', 'pdf2image', 'sympy']
for name in modules:
    importlib.import_module(name)
tools = ['pdftoppm', 'pdftotext', 'qpdf', 'sqlite3', 'unzip', 'zip', 'soffice', 'recalc-xlsx']
print(json.dumps({'missing_tools': [t for t in tools if shutil.which(t) is None]}))
""".strip()

    payload = _execute(client, code)

    assert json.loads(str(payload["stdout"])) == {"missing_tools": []}


def test_vietnamese_pdf_with_registered_font_round_trips() -> None:
    client = TestClient(create_app())
    code = f"""
from fpdf import FPDF
from pypdf import PdfReader
pdf = FPDF()
pdf.add_page()
pdf.add_font('DejaVu', fname='/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf')
pdf.set_font('DejaVu', size=14)
pdf.cell(text={VIETNAMESE!r})
pdf.output('report.pdf')
print(PdfReader('report.pdf').pages[0].extract_text().strip())
""".strip()

    payload = _execute(client, code)

    assert payload["stdout"] == VIETNAMESE + "\n"
    files = payload["files"]
    assert isinstance(files, list)
    file_id = next(f["file_id"] for f in files if f["path"] == "report.pdf")
    download = client.get(f"/v1/files/{file_id}")
    assert download.status_code == 200
    assert download.content.startswith(b"%PDF-")


def test_recalc_xlsx_fills_formula_values_and_reports_errors() -> None:
    client = TestClient(create_app())
    code = """
import json, subprocess, zipfile, openpyxl, xlsxwriter
from openpyxl.chart import BarChart, Reference
wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Doanh thu'
for row in [['Tháng', 'Doanh thu'], [1, 120], [2, 80], [3, 100]]:
    ws.append(row)
ws['B5'] = '=SUM(B2:B4)'; ws['B6'] = '=B5/0'
chart = BarChart()
data = Reference(ws, min_col=2, min_row=1, max_row=4)
chart.add_data(data, titles_from_data=True)
ws.add_chart(chart, 'D2'); wb.save('báo cáo.xlsx')
x = xlsxwriter.Workbook('cached.xlsx'); s = x.add_worksheet()
s.write_column('A1', [10, 20, 30]); s.write_formula('A4', '=SUM(A1:A3)'); x.close()
before = openpyxl.load_workbook('báo cáo.xlsx', data_only=True)
before = before['Doanh thu']['B5'].value
files = ['báo cáo.xlsx', 'cached.xlsx', 'macro.xlsm', 'cached.xlsx']
done = subprocess.run(['recalc-xlsx', *files], capture_output=True, text=True)
reports = [json.loads(line) for line in done.stdout.splitlines()]
book = openpyxl.load_workbook('báo cáo.xlsx', data_only=True)['Doanh thu']
names = zipfile.ZipFile('báo cáo.xlsx').namelist()
print(json.dumps({
    'returncode': done.returncode, 'before': before, 'sum': book['B5'].value,
    'cached': openpyxl.load_workbook('cached.xlsx', data_only=True).active['A4'].value,
    'formula': openpyxl.load_workbook('báo cáo.xlsx')['Doanh thu']['B5'].value,
    'charts': sum(n.startswith('xl/charts/chart') for n in names),
    'reports': reports,
}, ensure_ascii=False))
""".strip()

    response = client.post("/v1/execute", json={"code": code, "timeout_ms": 60000})
    assert response.status_code == 200
    payload = response.json()
    result = json.loads(str(payload["stdout"]))

    assert result["before"] is None
    assert result["sum"] == 300
    assert result["cached"] == 60
    assert result["formula"] == "=SUM(B2:B4)"
    assert result["charts"] == 1
    assert result["returncode"] == 1  # the .xlsm argument is refused
    assert result["reports"][0] == {
        "file": "báo cáo.xlsx",
        "recalculated": True,
        "error_count": 1,
        "errors": ["Doanh thu!B6: #DIV/0!"],
    }
    assert result["reports"][1]["recalculated"] is True
    assert result["reports"][2] == {
        "file": "macro.xlsm",
        "recalculated": False,
        "error": "only .xlsx files are supported",
    }
    assert result["reports"][3] == {
        "file": "cached.xlsx",
        "recalculated": False,
        "error": "run separately: same file name as another argument",
    }


def _files(payload: dict[str, object]) -> list[dict[str, Any]]:
    return cast(list[dict[str, Any]], payload["files"])


def test_open_figures_are_captured_as_chart_data_and_png_at_exit() -> None:
    client = TestClient(create_app())
    code = """
import matplotlib.pyplot as plt
fig, ax = plt.subplots()
ax.plot(['T1', 'T2', 'T3'], [120, 150, 90], label='Doanh thu')
ax.set_title('Doanh thu quý 3'); ax.set_ylabel('Tỷ đồng (₫)')
plt.figure(); plt.pie([40, 60], labels=['Hà Nội', 'Đà Nẵng'])
saved = plt.figure(); plt.bar(['A'], [1]); saved.savefig('mine.png'); plt.close(saved)
""".strip()

    payload = _execute(client, code)
    files = {entry["path"]: entry for entry in _files(payload)}
    assert {"mine.png", ".memoryos-charts/chart-1.png", ".memoryos-charts/chart-2.png"} <= set(
        files
    )
    assert ".memoryos-charts/chart-3.png" not in files  # a closed figure is not captured

    def read(path: str) -> bytes:
        response = client.get(f"/v1/files/{files[path]['file_id']}")
        assert response.status_code == 200
        return bytes(response.content)

    line = json.loads(read(".memoryos-charts/chart-1.json"))
    assert line["type"] == "line"
    assert line["title"] == "Doanh thu quý 3"
    assert line["y_unit"] == "₫"
    assert line["elements"] == [
        {"label": "Doanh thu", "points": [["T1", 120.0], ["T2", 150.0], ["T3", 90.0]]}
    ]
    pie = json.loads(read(".memoryos-charts/chart-2.json"))
    assert pie["type"] == "pie"
    assert [element["label"] for element in pie["elements"]] == ["Hà Nội", "Đà Nẵng"]
    assert read(".memoryos-charts/chart-1.png").startswith(b"\x89PNG")


def test_runs_without_pyplot_leave_no_chart_directory() -> None:
    payload = _execute(TestClient(create_app()), "print('xin chào')")
    assert not any(str(entry["path"]).startswith(".memoryos-charts") for entry in _files(payload))


def test_pptx_to_pdf_converts_a_vietnamese_deck_and_refuses_other_files() -> None:
    client = TestClient(create_app())
    code = """
import json, subprocess
from pptx import Presentation
deck = Presentation()
for number in range(3):
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = f'Doanh thu quý 3 — slide {number + 1}'
    slide.placeholders[1].text = 'Miền Bắc tăng 22,1% • Hà Nội, Đà Nẵng'
deck.save('báo cáo.pptx')
convert = ['pptx-to-pdf', 'báo cáo.pptx', 'preview.pdf']
done = subprocess.run(convert, capture_output=True, text=True)
text = subprocess.run(['pdftotext', 'preview.pdf', '-'], capture_output=True, text=True).stdout
refused = subprocess.run(['pptx-to-pdf', 'notes.txt', 'x.pdf'], capture_output=True, text=True)
result = {'code': done.returncode, 'report': json.loads(done.stdout), 'text': text,
          'refused': refused.returncode, 'refusal': json.loads(refused.stdout)}
print(json.dumps(result, ensure_ascii=False))
""".strip()

    response = client.post("/v1/execute", json={"code": code, "timeout_ms": 60000})
    assert response.status_code == 200
    result = json.loads(str(response.json()["stdout"]))

    assert result["code"] == 0
    assert result["report"] == {"converted": True, "pages": 3}
    assert "Doanh thu quý 3 — slide 2" in result["text"]
    assert "Miền Bắc tăng 22,1%" in result["text"]
    assert result["refused"] == 1
    assert result["refusal"] == {"converted": False, "error": "only .pptx files are supported"}


def test_files_written_to_mnt_data_are_returned() -> None:
    """ChatGPT-trained models save to /mnt/data; it is the workspace, so the file comes back."""
    client = TestClient(create_app())
    payload = _execute(
        client,
        "open('/mnt/data/bao-cao.txt', 'w', encoding='utf-8').write('Q3')\nprint('ok')",
    )

    files = cast(list[dict[str, Any]], payload["files"])
    assert [file["path"] for file in files] == ["bao-cao.txt"]
