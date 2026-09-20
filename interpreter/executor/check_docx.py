#!/opt/executor-venv/bin/python
"""Report the formatting faults that make a generated .docx look unformatted (MEM-110).

A model writing python-docx has no way to see its own output, so a document that uses
``add_paragraph`` for everything — no headings, borderless tables, bullets typed as text —
reads as one flat block and nobody notices until a person opens it. The Anthropic docx skill
answers the same gap by listing the footguns and rendering the result to look at it; the
sandbox cannot show a picture to the model, so this helper reports the same faults as text
the model can act on, as ``recalc-xlsx`` does for formula errors.

Usage: check-docx FILE.docx [FILE.docx ...]
Prints one JSON object per file: {"file", "paragraphs", "tables", "issue_count", "issues"}.
Each issue names what is wrong and how to fix it. No file is modified.
"""

from __future__ import annotations

import json
import re
import sys
from pathlib import Path

import docx
from docx.document import Document as DocxDocument
from docx.table import Table
from docx.text.paragraph import Paragraph

MAX_REPORTED_ISSUES = 20
# A short note needs no heading; a document long enough to scroll does.
HEADING_REQUIRED_PARAGRAPHS = 8
BULLET_MARKERS = ("•", "◦", "▪", "‣", "·", "- ", "* ", "+ ")
# python-docx creates tables with this style, which draws no lines at all.
BORDERLESS_TABLE_STYLES = {None, "", "Table Normal", "Normal Table", "TableNormal"}
# "1." or "2)" typed at the start of a paragraph, which Word cannot renumber.
SELF_NUMBERED = re.compile(r"^\d{1,3}[.)]\s")
LIST_STYLE_HINT = "List Bullet"
NUMBER_STYLE_HINT = "List Number"


def _style_name(item: Paragraph | Table) -> str:
    style = getattr(item, "style", None)
    name = getattr(style, "name", None)
    return name if isinstance(name, str) else ""


def _issues(document: DocxDocument) -> list[str]:
    paragraphs = [p for p in document.paragraphs if p.text.strip()]
    issues: list[str] = []

    headings = [p for p in paragraphs if _style_name(p).startswith(("Heading", "Title"))]
    if len(paragraphs) >= HEADING_REQUIRED_PARAGRAPHS and not headings:
        issues.append(
            f"no heading in {len(paragraphs)} paragraphs: give the document a title and sections "
            "with document.add_heading(text, level=0..3) so it has structure and a usable outline"
        )

    for index, paragraph in enumerate(paragraphs, start=1):
        text = paragraph.text.strip()
        style = _style_name(paragraph)
        if text.startswith(BULLET_MARKERS) and not style.startswith(("List", "ListParagraph")):
            issues.append(
                f"paragraph {index} starts with a typed bullet ({text[:20]!r}): drop the character "
                f"and use document.add_paragraph(text, style='{LIST_STYLE_HINT}'), or "
                f"'{NUMBER_STYLE_HINT}' for a numbered list"
            )
        elif SELF_NUMBERED.match(text) and not style.startswith("List"):
            issues.append(
                f"paragraph {index} numbers itself ({text[:20]!r}): use "
                f"document.add_paragraph(text, style='{NUMBER_STYLE_HINT}') so Word renumbers it"
            )
        if "\n" in paragraph.text or "\r" in paragraph.text:
            issues.append(
                f"paragraph {index} contains a line break character: Word keeps one paragraph, so "
                "call add_paragraph once per paragraph instead"
            )

    for index, table in enumerate(document.tables, start=1):
        if _style_name(table) in BORDERLESS_TABLE_STYLES:
            issues.append(
                f"table {index} has no borders: set table.style = 'Table Grid' so its cells are "
                "readable"
            )
        header = table.rows[0].cells if table.rows else []
        if header and all(not cell.text.strip() for cell in header):
            issues.append(f"table {index} has an empty first row: put the column names there")

    if not paragraphs and not document.tables:
        issues.append("the document is empty: nothing was written before saving")
    return issues


def _report(path: Path) -> dict[str, object]:
    if path.suffix.lower() != ".docx":
        return {"file": path.name, "checked": False, "error": "only .docx files are supported"}
    if not path.is_file():
        return {"file": path.name, "checked": False, "error": "file not found"}
    try:
        document = docx.Document(str(path))
    except Exception as failure:  # noqa: BLE001 - any unreadable file is reported, never raised
        return {
            "file": path.name,
            "checked": False,
            "error": f"{type(failure).__name__}: {failure}"[:200],
        }
    issues = _issues(document)
    return {
        "file": path.name,
        "checked": True,
        "paragraphs": len([p for p in document.paragraphs if p.text.strip()]),
        "tables": len(document.tables),
        "issue_count": len(issues),
        "issues": issues[:MAX_REPORTED_ISSUES],
    }


def main(argv: list[str]) -> int:
    if not argv:
        print("usage: check-docx FILE.docx [FILE.docx ...]", file=sys.stderr)
        return 2
    reports = [_report(Path(argument)) for argument in argv]
    for report in reports:
        print(json.dumps(report, ensure_ascii=False))
    # A fault is a report, not a failed command: the model reads the JSON and decides what to fix.
    return 0 if all(report.get("checked") for report in reports) else 1


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
