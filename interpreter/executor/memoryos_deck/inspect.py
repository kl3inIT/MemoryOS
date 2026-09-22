"""The slide faults a model cannot see in the deck it just wrote (MEM-122).

``check-docx`` exists because a model writing python-docx cannot look at its own document. A deck is worse:
a flat Word document is ugly but readable, while a slide whose text runs past the edge is unusable and the
model has no way to notice. These checks read a finished ``.pptx`` — one this package rendered or one the
model built by hand with python-pptx — and report the faults as text it can act on.

Text height is estimated with :mod:`memoryos_deck.measure`, the same estimator the renderer fits with, so a
deck produced by ``render-deck`` reports no overflow.
"""

from __future__ import annotations

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn
from pptx.util import Pt

from . import measure, theme
from .render import MARKER

MAX_REPORTED_ISSUES = 25
#: Nothing smaller than a heading on a slide means the slide has no title.
MIN_TITLE_PT = 20
#: python-pptx leaves no explicit size on an inherited run; the Office body default is 18 pt.
DEFAULT_BODY_PT = 18
MAX_BULLETS = 7
LONG_BULLET_CHARS = 200
#: The estimator is approximate, so only a real overflow is reported.
OVERFLOW_TOLERANCE = 1.06
PLACEHOLDER_MARKERS = ("lorem ipsum", "click to edit", "placeholder text", "your text here",
                       "tbd", "todo", "xxx", "[insert", "<insert")
#: A picture scaled off its native aspect ratio is squashed; the renderer never does this.
ASPECT_TOLERANCE = 0.03
#: A line the audience reads as a bullet: a real PowerPoint bullet, or one typed as a character.
TYPED_BULLETS = ("•", "◦", "▪", "‣", "·", "- ", "– ", "* ")


def _runs(frame) -> list:
    return [run for paragraph in frame.paragraphs for run in paragraph.runs]


def _is_bullet(paragraph) -> bool:
    properties = paragraph._p.pPr
    if properties is not None and properties.find(qn("a:buChar")) is not None:
        return True
    return "".join(run.text for run in paragraph.runs).strip().startswith(TYPED_BULLETS)


def _size_pt(run) -> float | None:
    size = run.font.size
    return None if size is None else size.pt


def _paragraph_pt(paragraph) -> float:
    sizes = [size for size in (_size_pt(run) for run in paragraph.runs) if size]
    return max(sizes) if sizes else DEFAULT_BODY_PT


def _needed_inches(frame, width_in: float) -> float:
    total = 0.0
    for paragraph in frame.paragraphs:
        text = "".join(run.text for run in paragraph.runs)
        size = _paragraph_pt(paragraph)
        lines = max(1, len(measure.wrap(text, width_in, size))) if text.strip() else 1
        spacing = paragraph.line_spacing if isinstance(paragraph.line_spacing, float) else measure.LINE_SPACING
        total += lines * size * spacing / 72.0
        after = paragraph.space_after
        if after is not None:
            total += after.pt / 72.0
    return total


def _frame_width_inches(shape) -> float:
    frame = shape.text_frame
    margins = (frame.margin_left or Pt(0)).inches + (frame.margin_right or Pt(0)).inches
    return max(0.1, shape.width.inches - margins)


def _text_shapes(slide) -> list:
    return [shape for shape in slide.shapes
            if shape.has_text_frame and shape.text_frame.text.strip()]


def _slide_issues(slide, number: int, width_in: float, height_in: float) -> list[str]:
    issues: list[str] = []
    where = f"slide {number}"
    shapes = _text_shapes(slide)
    tables = [shape for shape in slide.shapes if getattr(shape, "has_table", False)]
    pictures = [shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]

    if not shapes and not tables and not pictures:
        issues.append(f"{where} is empty: it carries no text, table or picture")
        return issues

    sizes = [size for shape in shapes for size in (_size_pt(run) for run in _runs(shape.text_frame)) if size]
    if shapes and sizes and max(sizes) < MIN_TITLE_PT:
        issues.append(f"{where} has no text larger than {MIN_TITLE_PT} pt: give it a title the audience "
                      "can read from the back of the room")
    for size in sorted({size for size in sizes if size < theme.MIN_READABLE_PT}):
        issues.append(f"{where} contains {size:g} pt text: nothing under {theme.MIN_READABLE_PT} pt is "
                      "readable when projected — move the content to another slide instead of shrinking it")

    for shape in shapes:
        frame = shape.text_frame
        text = frame.text.strip()
        lowered = text.lower()
        for marker in PLACEHOLDER_MARKERS:
            if marker in lowered:
                issues.append(f"{where} still contains placeholder text ({text[:40]!r}): replace it or "
                              "remove the shape")
                break
        needed = _needed_inches(frame, _frame_width_inches(shape))
        if needed > shape.height.inches * OVERFLOW_TOLERANCE:
            issues.append(f"{where}: text needs about {needed:.2f} in but its box is "
                          f"{shape.height.inches:.2f} in — the text runs past the edge; shorten it or "
                          "split the slide")
        lines = [paragraph for paragraph in frame.paragraphs
                 if "".join(run.text for run in paragraph.runs).strip()]
        bullets = [paragraph for paragraph in lines if _is_bullet(paragraph)]
        if len(bullets) > MAX_BULLETS:
            issues.append(f"{where} has {len(bullets)} bullets in one box; at most {MAX_BULLETS} belong on "
                          "a slide — split it across slides")
        for paragraph in lines:
            line = "".join(run.text for run in paragraph.runs).strip()
            if len(line) > LONG_BULLET_CHARS:
                issues.append(f"{where}: a line is {len(line)} characters ({line[:40]!r}…); a slide line is "
                              "a claim, not a paragraph")
                break

    for shape in slide.shapes:
        if shape.left is None or shape.top is None or shape.width is None or shape.height is None:
            continue
        right = shape.left.inches + shape.width.inches
        bottom = shape.top.inches + shape.height.inches
        if shape.left.inches < -0.01 or shape.top.inches < -0.01 or right > width_in + 0.01 \
                or bottom > height_in + 0.01:
            issues.append(f"{where}: a shape lies partly outside the slide "
                          f"({shape.left.inches:.1f}, {shape.top.inches:.1f} to {right:.1f}, {bottom:.1f} "
                          f"in on a {width_in:.1f} x {height_in:.1f} in slide)")

    for index, picture in enumerate(pictures, start=1):
        native = getattr(picture.image, "size", None)
        if not native or not native[0] or not native[1] or not picture.height:
            continue
        drawn = picture.width.inches / picture.height.inches
        original = native[0] / native[1]
        if abs(drawn - original) / original > ASPECT_TOLERANCE:
            issues.append(f"{where}: picture {index} is stretched (drawn {drawn:.2f}:1, original "
                          f"{original:.2f}:1); scale both sides by the same factor")
    return issues


def report(path: str, name: str) -> dict[str, object]:
    """One JSON-ready report for a finished deck, in the shape ``check-docx`` uses."""
    try:
        presentation = Presentation(path)
    except Exception as failure:  # noqa: BLE001 - an unreadable deck is reported, never raised
        return {"file": name, "checked": False, "error": f"{type(failure).__name__}: {failure}"[:200]}

    width_in = presentation.slide_width / 914400
    height_in = presentation.slide_height / 914400
    issues: list[str] = []
    slides = list(presentation.slides)
    if not slides:
        issues.append("the deck has no slides: nothing was added before saving")
    for number, slide in enumerate(slides, start=1):
        issues.extend(_slide_issues(slide, number, width_in, height_in))

    rendered = presentation.core_properties.category == MARKER
    if not rendered:
        issues.append("this deck was not built by render-deck: write a deck plan and render it so the "
                      "geometry, type sizes and footers are the same on every slide")
    return {"file": name, "checked": True, "slides": len(slides), "built_with_render_deck": rendered,
            "issue_count": len(issues), "issues": issues[:MAX_REPORTED_ISSUES]}
